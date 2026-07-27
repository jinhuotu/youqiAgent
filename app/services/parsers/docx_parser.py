"""Word(DOCX) 文本抽取。"""

from pathlib import Path

from docx import Document

from app.core.exceptions import BusinessError
from app.services.parsers.cleaning import clean_extracted_text


def extract_docx_text(path: Path) -> str:
    """从 DOCX 抽取段落与表格文本。"""
    try:
        doc = Document(str(path))
        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = clean_extracted_text("\n".join(parts))
        if not text:
            raise BusinessError("Word 文档未提取到文本")
        return text
    except BusinessError:
        raise
    except Exception as e:
        raise BusinessError(f"Word 解析失败: {e}") from e
