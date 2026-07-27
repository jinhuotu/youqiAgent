"""Excel(XLSX) 文本抽取。"""

from pathlib import Path

from openpyxl import load_workbook

from app.core.exceptions import BusinessError
from app.services.parsers.cleaning import clean_extracted_text


def extract_xlsx_text(path: Path) -> str:
    """从 XLSX 按工作表抽取行文本。"""
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            parts.append(f"## 工作表: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
            parts.append("")
        wb.close()
        text = clean_extracted_text("\n".join(parts))
        if not text:
            raise BusinessError("Excel 未提取到文本")
        return text
    except BusinessError:
        raise
    except Exception as e:
        raise BusinessError(f"Excel 解析失败: {e}") from e
